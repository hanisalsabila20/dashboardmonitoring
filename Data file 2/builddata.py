"""
Membuat data.js dari file Excel "Dashboard - HTML.xlsx".

Cara pakai (butuh Python + openpyxl):
    pip install openpyxl
    python build_data.py "Dashboard - HTML.xlsx"

Hasilnya: data.js (menimpa file lama). Jangan ubah index.html / style.css / app.js,
cukup refresh browser setelah data.js dibuat ulang.
"""
import json
import sys
import openpyxl

XLSX = sys.argv[1] if len(sys.argv) > 1 else 'Dashboard - HTML.xlsx'
OUT = 'data.js'

# Kolom sheet "KPI DSE" -> nama field (urutan = nomor kolom, mulai 0)
# None = kolom dilewati
DSE_COLS = [
    None, 'nom', 'site', 'id', 'cso', 'kpi', 'ci', 'comp',
    'osaT', 'osa', 'osaGap', 'osaIm3', 'osa3id', 'osaPct', 'osaKpi', 'desc',
    'spT', 'sp', 'spGap', 'spIm3', 'sp3id', 'spPct', 'spKpi',
    'ecT', 'ec', 'ecGap', 'ecAch', 'ecComp',
    'rrT', 'rr', 'rrGap', 'rrAch', 'rrComp',
    'lrsT', 'lrs', 'lrsL', 'lrsGap', 'lrsBase',
    'gaT', 'ga', 'gaGap', None,
    'pT', 'p', 'pGap', 'pAch',
    'o3T', 'o3', 'o3Gap', 'o3Ach',
    'tgT', 'tg', 'tgGap', 'tgAch',
]

# Kolom sheet "Outlet Wise" yang dipakai (urutan = urutan kolom di tabel dashboard)
# DSE, Site, PJP, Outlet, ID IM3, ID 3ID, Target, Realisasi, Gap, Ach, IM3, 3ID,
# Status, Total, IM3, 3ID (Sellin SP), Status, Total, IM3, 3ID (Retail tagging)
OUTLET_COLS = [1, 3, 4, 7, 5, 6, 8, 9, 10, 11, 12, 13, 15, 16, 19, 20, 21, 22, 23, 24]


def clean(v):
    """Angka bulat disimpan sebagai int, string ID tetap string."""
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


wb = openpyxl.load_workbook(XLSX, data_only=True)

# ---- Sheet KPI DSE ----
ws = wb['KPI DSE']
date = None
dse = []
for row in ws.iter_rows(values_only=True):
    if date is None and hasattr(row[7], 'strftime'):
        date = row[7].strftime('%Y-%m-%d')
    if isinstance(row[3], str) and row[3].startswith('MUJABEAST'):
        item = {}
        for i, key in enumerate(DSE_COLS):
            if key:
                item[key] = clean(row[i])
        dse.append(item)

# ---- Sheet Outlet Wise ----
ws = wb['Outlet Wise']
outlets = []
for row in ws.iter_rows(values_only=True):
    if isinstance(row[1], str) and row[1].startswith('MUJABEAST'):
        item = []
        for n, c in enumerate(OUTLET_COLS):
            v = clean(row[c])
            if n < 6:                      # 6 kolom pertama = teks
                v = '' if v is None else str(v)
            item.append(v)
        outlets.append(item)

# ---- Tulis data.js ----
lines = [
    '// Data dashboard - dibuat otomatis oleh build_data.py dari file Excel',
    '// Jangan edit manual; ubah Excel lalu jalankan ulang build_data.py',
    'const D = {',
    f'  date: {json.dumps(date)},',
    '  dse: [',
]
for i, d in enumerate(dse):
    lines.append('    {')
    keys = list(d)
    for j, k in enumerate(keys):
        comma = ',' if j < len(keys) - 1 else ''
        lines.append(f'      {k}: {json.dumps(d[k], ensure_ascii=False)}{comma}')
    lines.append('    }' + (',' if i < len(dse) - 1 else ''))
lines += [
    '  ],',
    '  // Kolom: DSE, Site, PJP, Outlet, ID IM3, ID 3ID, Target, Realisasi, Gap, Ach, IM3, 3ID,',
    '  //        Status OSA, Sellin SP (Total, IM3, 3ID), Status, Retail tagging (Total, IM3, 3ID)',
    '  outlets: [',
]
for i, r in enumerate(outlets):
    lines.append('    ' + json.dumps(r, ensure_ascii=False) + (',' if i < len(outlets) - 1 else ''))
lines += ['  ]', '};']

with open(OUT, 'w', encoding='utf-8') as f:
    f.write('\n'.join(lines) + '\n')

print(f'OK: {len(dse)} DSE, {len(outlets)} outlet, tanggal {date} -> {OUT}')